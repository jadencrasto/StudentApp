import io
import re
import string
import secrets
from typing import Any
import uuid
from uuid import UUID

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from app.models.user import User
from app.models.institution import Cohort
from app.core.security import hash_password


# RFC 5322 simplified email regex
EMAIL_REGEX = re.compile(r"^[\w\.\+\-]+@[a-zA-Z0-9\.\-]+\.[a-zA-Z]{2,}$")


def generate_temporary_password(length: int = 12) -> str:
    """
    Generates a strong temporary password (>= 12 characters).
    CRITICAL REQUIREMENT: Must ALWAYS start with an ASCII letter.
    Contains at least one uppercase, one lowercase, one digit, and one special character.
    """
    if length < 12:
        length = 12

    # Guaranteed to start with an ASCII letter
    first_char = secrets.choice(string.ascii_letters)

    specials = "!@#$%^&*"
    mandatory = [
        secrets.choice(string.ascii_uppercase),
        secrets.choice(string.ascii_lowercase),
        secrets.choice(string.digits),
        secrets.choice(specials),
    ]

    all_pool = string.ascii_letters + string.digits + specials
    remaining_count = length - 1 - len(mandatory)
    fillers = [secrets.choice(all_pool) for _ in range(remaining_count)]

    tail = mandatory + fillers
    secrets.SystemRandom().shuffle(tail)

    return first_char + "".join(tail)


async def generate_unique_username(
    full_name: str,
    db: AsyncSession,
    reserved_usernames: set[str] | None = None,
) -> str:
    """
    Generates a clean, lowercase alphanumeric username guaranteed unique across the whole DB
    and any in-memory reserved set.
    """
    reserved = reserved_usernames or set()

    # Clean full name to lowercase alphanumeric
    base = re.sub(r"[^a-zA-Z0-9]", "", full_name.lower().strip())
    if len(base) < 3:
        base = "student"
    base = base[:15]  # limit base length

    for _ in range(100):
        # 4-digit random suffix
        suffix = f"{secrets.randbelow(9000) + 1000}"
        candidate = f"{base}{suffix}"

        if candidate in reserved:
            continue

        existing = await db.scalar(select(User.id).where(User.username == candidate))
        if existing is None:
            reserved.add(candidate)
            return candidate

    # Fallback to UUID prefix if collisions persist
    fallback = f"{base}_{uuid.uuid4().hex[:6]}"
    reserved.add(fallback)
    return fallback


def sanitize_text_cell(value: Any) -> str:
    """
    Spreadsheet injection protection:
    If a text cell starts with '=', '+', '-', or '@', prefix it with "'"
    so spreadsheet engines treat it as raw text rather than an executable formula.
    """
    if value is None:
        return ""
    s = str(value).strip()
    if s and s[0] in ("=", "+", "-", "@"):
        return f"'{s}"
    return s


def generate_student_template_xlsx() -> bytes:
    """
    Generates a clean .xlsx template for bulk student upload.
    Headers: Full Name, Email, Class
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students Template"

    # Header styling
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")  # Emerald 600
    border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )

    headers = ["Full Name", "Email", "Class"]
    ws.append(headers)

    for col_idx in range(1, 4):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    # Sample rows
    samples = [
        ["Aarav Sharma", "aarav.sharma@example.edu", "Computer Science - Year 2"],
        ["Diya Patel", "diya.patel@example.edu", "Computer Science - Year 2"],
        ["Rohan Gupta", "rohan.gupta@example.edu", "Information Technology - Year 1"],
    ]

    for row_data in samples:
        ws.append(row_data)

    for row_idx in range(2, len(samples) + 2):
        for col_idx in range(1, 4):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = border
            cell.number_format = '@'

    # Auto-adjust column widths
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def parse_and_preview_students_xlsx(
    file_bytes: bytes,
    institution_id: UUID,
    db: AsyncSession,
) -> dict[str, Any]:
    """
    Parses an uploaded .xlsx file and generates a validation preview WITHOUT saving.
    Enforces <= 2MB size, <= 500 rows, valid email format, uniqueness, and institution cohort matching.
    """
    if len(file_bytes) > 2 * 1024 * 1024:
        raise ValueError("File size exceeds 2 MB limit")

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel file is empty")

    # Find headers (row 1)
    header_row = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    
    col_name_idx = -1
    col_email_idx = -1
    col_class_idx = -1

    for idx, header in enumerate(header_row):
        h = header.replace("_", " ").replace("-", " ")
        if "name" in h and "full" in h or h == "name" or "student name" in h:
            col_name_idx = idx
        elif "email" in h:
            col_email_idx = idx
        elif "class" in h or "cohort" in h or "grade" in h:
            col_class_idx = idx

    if col_name_idx == -1 or col_email_idx == -1 or col_class_idx == -1:
        raise ValueError(
            "Excel file must include 'Full Name', 'Email', and 'Class' (or Cohort) columns in the header row"
        )

    data_rows = rows[1:]
    # Filter out completely blank rows
    non_empty_rows = []
    for r_idx, r in enumerate(data_rows, start=2):
        if any(cell is not None and str(cell).strip() != "" for cell in r):
            non_empty_rows.append((r_idx, r))

    if len(non_empty_rows) > 500:
        raise ValueError(f"Upload exceeds maximum limit of 500 rows (found {len(non_empty_rows)} rows)")

    if len(non_empty_rows) == 0:
        raise ValueError("No student records found in file")

    # Fetch active cohorts for this institution
    cohort_stmt = select(Cohort).where(
        Cohort.institution_id == institution_id,
        Cohort.is_active == True,
    )
    res = await db.execute(cohort_stmt)
    active_cohorts = res.scalars().all()

    # Build cohort lookup: case-insensitive name -> Cohort
    cohort_lookup = {c.name.strip().lower(): c for c in active_cohorts}

    # Extract all emails to query existing in DB in bulk
    all_emails = []
    for _, r in non_empty_rows:
        if len(r) > col_email_idx and r[col_email_idx]:
            em = str(r[col_email_idx]).strip().lower()
            all_emails.append(em)

    existing_emails_in_db = set()
    if all_emails:
        db_emails_stmt = select(User.email).where(User.email.in_(all_emails))
        db_res = await db.execute(db_emails_stmt)
        existing_emails_in_db = {row[0].lower() for row in db_res.all()}

    valid_rows = []
    invalid_rows = []
    seen_file_emails = set()

    for row_num, r in non_empty_rows:
        raw_name = r[col_name_idx] if len(r) > col_name_idx else None
        raw_email = r[col_email_idx] if len(r) > col_email_idx else None
        raw_class = r[col_class_idx] if len(r) > col_class_idx else None

        name_str = str(raw_name).strip() if raw_name is not None else ""
        email_str = str(raw_email).strip().lower() if raw_email is not None else ""
        class_str = str(raw_class).strip() if raw_class is not None else ""

        reasons = []

        if not name_str:
            reasons.append("Missing full name")
        if not email_str:
            reasons.append("Missing email")
        elif not EMAIL_REGEX.match(email_str):
            reasons.append("Invalid email format")
        elif email_str in seen_file_emails:
            reasons.append("Duplicate email within uploaded file")
        elif email_str in existing_emails_in_db:
            reasons.append("Email already registered in system")

        cohort_obj = None
        if not class_str:
            reasons.append("Missing class")
        else:
            cohort_obj = cohort_lookup.get(class_str.lower())
            if not cohort_obj:
                reasons.append(f"Class '{class_str}' does not exist in your institution")

        if email_str:
            seen_file_emails.add(email_str)

        row_item = {
            "row_num": row_num,
            "full_name": name_str,
            "email": email_str,
            "class_name": class_str,
            "cohort_id": str(cohort_obj.id) if cohort_obj else None,
            "valid": len(reasons) == 0,
            "reason": "; ".join(reasons) if reasons else None,
        }

        if len(reasons) == 0:
            valid_rows.append(row_item)
        else:
            invalid_rows.append(row_item)

    return {
        "total_rows": len(non_empty_rows),
        "valid_count": len(valid_rows),
        "invalid_count": len(invalid_rows),
        "valid_rows": valid_rows,
        "invalid_rows": invalid_rows,
    }


def generate_bad_rows_xlsx(bad_rows: list[dict[str, Any]]) -> bytes:
    """
    Generates a downloadable Excel report of invalid rows with reasons.
    Applies spreadsheet injection protection to names, emails, and reasons.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Upload Errors"

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")  # Red 600
    border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )

    headers = ["Row #", "Full Name", "Email", "Class", "Error Reason"]
    ws.append(headers)

    for col_idx in range(1, 6):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row_item in bad_rows:
        ws.append([
            row_item.get("row_num", ""),
            sanitize_text_cell(row_item.get("full_name", "")),
            sanitize_text_cell(row_item.get("email", "")),
            sanitize_text_cell(row_item.get("class_name", "")),
            sanitize_text_cell(row_item.get("reason", "")),
        ])

    for row_idx in range(2, len(bad_rows) + 2):
        for col_idx in range(1, 6):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = border
            cell.number_format = '@'

    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 35
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 45

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_credentials_xlsx(credentials: list[dict[str, Any]]) -> bytes:
    """
    Builds the ONE-TIME downloadable credentials workbook (.xlsx) strictly in-memory.
    Requirement:
      - Plain-text passwords written as explicit text cells (data_type='s') without leading apostrophes.
      - User-supplied names and emails protected against spreadsheet formula injection.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Credentials"

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")  # Emerald 600
    border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )

    headers = ["Full Name", "Email", "Username", "Temporary Password", "Class"]
    ws.append(headers)

    for col_idx in range(1, 6):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for item in credentials:
        ws.append([
            sanitize_text_cell(item.get("full_name", "")),
            sanitize_text_cell(item.get("email", "")),
            str(item.get("username", "")),
            str(item.get("temporary_password", "")),  # Explicit text, no apostrophe
            sanitize_text_cell(item.get("class_name", "")),
        ])

    for row_idx in range(2, len(credentials) + 2):
        for col_idx in range(1, 6):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = border
            # Explicit text format so passwords don't trigger numeric or scientific format
            cell.number_format = '@'
            cell.data_type = 's'

    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
