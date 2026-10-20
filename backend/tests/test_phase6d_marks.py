import io
import uuid
from datetime import date
from decimal import Decimal
import pytest
import openpyxl
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func

from app.models.user import User
from app.models.institution import Institution, Cohort, generate_join_code
from app.models.class_marks import Exam, Mark, MarkChangeRecord
from app.core.security import hash_password


async def create_institution(
    db: AsyncSession,
    name: str = "Test College",
    inst_type: str = "college",
) -> Institution:
    inst = Institution(
        name=name,
        type=inst_type,
        location="Mumbai",
        join_code=generate_join_code(),
        is_active=True,
    )
    db.add(inst)
    await db.commit()
    await db.refresh(inst)
    return inst


async def create_cohort(
    db: AsyncSession,
    institution_id: uuid.UUID,
    name: str = "Class A",
    level_type: str = "grade",
    level_value: str = "10",
    academic_year: str = "2026-27",
) -> Cohort:
    cohort = Cohort(
        institution_id=institution_id,
        name=name,
        level_type=level_type,
        level_value=level_value,
        academic_year=academic_year,
        is_active=True,
    )
    db.add(cohort)
    await db.commit()
    await db.refresh(cohort)
    return cohort


async def create_user(
    db: AsyncSession,
    email: str,
    role: str = "student",
    institution_id: uuid.UUID | None = None,
    cohort_id: uuid.UUID | None = None,
    is_active: bool = True,
    password: str = "Password123!",
) -> User:
    user = User(
        email=email,
        username=email.split("@")[0] + "_" + uuid.uuid4().hex[:4],
        full_name=f"User {email}",
        hashed_password=hash_password(password),
        role=role,
        institution_id=institution_id,
        cohort_id=cohort_id,
        is_active=is_active,
        token_version=1,
    )
    db.add(user)
    await db.commit()
    await db.refresh(user)
    return user


async def login(client: AsyncClient, email: str, password: str = "Password123!") -> str:
    resp = await client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert resp.status_code == 200, resp.text
    return resp.json()["access_token"]


@pytest.mark.asyncio
class TestPhase6DClassMarks:

    async def test_cross_institution_blocked(self, client: AsyncClient, db_session: AsyncSession):
        inst1 = await create_institution(db_session, "Inst 1")
        inst2 = await create_institution(db_session, "Inst 2")
        cohort2 = await create_cohort(db_session, inst2.id, "Cohort 2")

        admin1 = await create_user(db_session, f"admin_{uuid.uuid4().hex[:6]}@inst1.edu", role="admin", institution_id=inst1.id)
        token1 = await login(client, admin1.email)
        headers1 = {"Authorization": f"Bearer {token1}"}

        # Admin 1 attempts to create exam in Inst 2 cohort -> 404
        resp = await client.post(
            "/api/v1/admin/marks/exams",
            json={
                "cohort_id": str(cohort2.id),
                "exam_type": "internals",
                "title": "Unit Test 1",
                "subject": "Physics",
                "max_marks": 25.0,
            },
            headers=headers1,
        )
        assert resp.status_code == 404
        assert "not found in your institution" in resp.json()["detail"].lower()

    async def test_duplicate_exam_blocked(self, client: AsyncClient, db_session: AsyncSession):
        inst = await create_institution(db_session, "Dup Inst")
        cohort = await create_cohort(db_session, inst.id, "Class 10")
        admin = await create_user(db_session, f"admin_{uuid.uuid4().hex[:6]}@inst.edu", role="admin", institution_id=inst.id)
        token = await login(client, admin.email)
        headers = {"Authorization": f"Bearer {token}"}

        # Create first exam
        resp1 = await client.post(
            "/api/v1/admin/marks/exams",
            json={
                "cohort_id": str(cohort.id),
                "exam_type": "mid-term",
                "title": "  Midterm Exam  ",
                "subject": "  Chemistry  ",
                "max_marks": 50.0,
            },
            headers=headers,
        )
        assert resp1.status_code == 200, resp1.text

        # Create duplicate exam with different casing & spacing -> 400
        resp2 = await client.post(
            "/api/v1/admin/marks/exams",
            json={
                "cohort_id": str(cohort.id),
                "exam_type": "MID-TERM",
                "title": "midterm exam",
                "subject": "chemistry",
                "max_marks": 50.0,
            },
            headers=headers,
        )
        assert resp2.status_code == 400
        assert "already exists" in resp2.json()["detail"].lower()

    async def test_marks_above_max_rejected(self, client: AsyncClient, db_session: AsyncSession):
        inst = await create_institution(db_session, "Max Marks Inst")
        cohort = await create_cohort(db_session, inst.id, "Grade 9")
        admin = await create_user(db_session, f"admin_{uuid.uuid4().hex[:6]}@inst.edu", role="admin", institution_id=inst.id)
        student = await create_user(db_session, f"st_{uuid.uuid4().hex[:6]}@inst.edu", role="student", institution_id=inst.id, cohort_id=cohort.id)

        token = await login(client, admin.email)
        headers = {"Authorization": f"Bearer {token}"}

        resp = await client.post(
            "/api/v1/admin/marks/exams",
            json={
                "cohort_id": str(cohort.id),
                "exam_type": "quiz",
                "title": "Quiz 1",
                "subject": "Biology",
                "max_marks": 20.0,
            },
            headers=headers,
        )
        exam_id = resp.json()["id"]

        # Attempt to save 25 marks on a 20 mark exam -> 400
        save_resp = await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/marks",
            json={
                "marks": [
                    {"student_id": str(student.id), "marks_obtained": 25.0, "is_absent": False}
                ]
            },
            headers=headers,
        )
        assert save_resp.status_code == 400
        assert "exceed maximum" in save_resp.json()["detail"].lower()

    async def test_max_marks_cannot_be_lowered_below_existing_mark(self, client: AsyncClient, db_session: AsyncSession):
        inst = await create_institution(db_session, "Lower Max Inst")
        cohort = await create_cohort(db_session, inst.id, "Grade 11")
        admin = await create_user(db_session, f"admin_{uuid.uuid4().hex[:6]}@inst.edu", role="admin", institution_id=inst.id)
        student = await create_user(db_session, f"st_{uuid.uuid4().hex[:6]}@inst.edu", role="student", institution_id=inst.id, cohort_id=cohort.id)

        token = await login(client, admin.email)
        headers = {"Authorization": f"Bearer {token}"}

        resp = await client.post(
            "/api/v1/admin/marks/exams",
            json={
                "cohort_id": str(cohort.id),
                "exam_type": "internals",
                "title": "Internal 1",
                "subject": "Math",
                "max_marks": 30.0,
            },
            headers=headers,
        )
        exam_id = resp.json()["id"]

        # Enter mark of 28.5
        save_resp = await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/marks",
            json={
                "marks": [
                    {"student_id": str(student.id), "marks_obtained": 28.5, "is_absent": False}
                ]
            },
            headers=headers,
        )
        assert save_resp.status_code == 200

        # Attempt to patch max_marks to 25 -> 400
        patch_resp = await client.patch(
            f"/api/v1/admin/marks/exams/{exam_id}",
            json={"max_marks": 25.0},
            headers=headers,
        )
        assert patch_resp.status_code == 400
        assert "cannot reduce max marks" in patch_resp.json()["detail"].lower()

    async def test_absent_differs_from_zero(self, client: AsyncClient, db_session: AsyncSession):
        inst = await create_institution(db_session, "Absent Inst")
        cohort = await create_cohort(db_session, inst.id, "Division B")
        admin = await create_user(db_session, f"admin_{uuid.uuid4().hex[:6]}@inst.edu", role="admin", institution_id=inst.id)
        st_absent = await create_user(db_session, f"absent_{uuid.uuid4().hex[:6]}@inst.edu", role="student", institution_id=inst.id, cohort_id=cohort.id)
        st_zero = await create_user(db_session, f"zero_{uuid.uuid4().hex[:6]}@inst.edu", role="student", institution_id=inst.id, cohort_id=cohort.id)

        token = await login(client, admin.email)
        headers = {"Authorization": f"Bearer {token}"}

        resp = await client.post(
            "/api/v1/admin/marks/exams",
            json={
                "cohort_id": str(cohort.id),
                "exam_type": "end-term",
                "title": "Semester Exam",
                "subject": "English",
                "max_marks": 100.0,
            },
            headers=headers,
        )
        exam_id = resp.json()["id"]

        # Save absent for st_absent and 0.0 for st_zero
        save_resp = await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/marks",
            json={
                "marks": [
                    {"student_id": str(st_absent.id), "is_absent": True, "marks_obtained": None},
                    {"student_id": str(st_zero.id), "is_absent": False, "marks_obtained": 0.0},
                ]
            },
            headers=headers,
        )
        assert save_resp.status_code == 200

        # Verify DB directly
        m_abs = await db_session.scalar(select(Mark).where(Mark.exam_id == uuid.UUID(exam_id), Mark.student_id == st_absent.id))
        m_zer = await db_session.scalar(select(Mark).where(Mark.exam_id == uuid.UUID(exam_id), Mark.student_id == st_zero.id))

        assert m_abs.is_absent is True
        assert m_abs.marks_obtained is None  # NOT 0!

        assert m_zer.is_absent is False
        assert m_zer.marks_obtained == Decimal("0.00")

    async def test_student_with_no_mark_sees_not_entered_yet_never_zero(self, client: AsyncClient, db_session: AsyncSession):
        inst = await create_institution(db_session, "View Inst")
        cohort = await create_cohort(db_session, inst.id, "Class 12")
        admin = await create_user(db_session, f"admin_{uuid.uuid4().hex[:6]}@inst.edu", role="admin", institution_id=inst.id)
        student = await create_user(db_session, f"st_{uuid.uuid4().hex[:6]}@inst.edu", role="student", institution_id=inst.id, cohort_id=cohort.id)

        admin_token = await login(client, admin.email)
        student_token = await login(client, student.email)

        # Admin creates and publishes exam without entering mark for student
        e_resp = await client.post(
            "/api/v1/admin/marks/exams",
            json={
                "cohort_id": str(cohort.id),
                "exam_type": "unit-test",
                "title": "Unit Test 1",
                "subject": "History",
                "max_marks": 25.0,
            },
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        exam_id = e_resp.json()["id"]

        pub_resp = await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/publish",
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        assert pub_resp.status_code == 200
        # Check warning about unentered marks
        assert len(pub_resp.json()["warnings"]) > 0
        assert "no marks entered" in pub_resp.json()["warnings"][0].lower()

        # Student views marks
        st_resp = await client.get("/api/v1/student/marks", headers={"Authorization": f"Bearer {student_token}"})
        assert st_resp.status_code == 200
        data = st_resp.json()
        assert data["enrolled"] is True
        assert len(data["exams"]) == 1

        ex_item = data["exams"][0]
        assert ex_item["status"] == "not_entered"
        assert ex_item["marks_obtained"] is None
        assert ex_item["percentage"] is None
        assert ex_item["display_text"] == "Not entered yet"

    async def test_draft_hidden_and_unpublish_hides_again(self, client: AsyncClient, db_session: AsyncSession):
        inst = await create_institution(db_session, "Draft Inst")
        cohort = await create_cohort(db_session, inst.id, "Grade 8")
        admin = await create_user(db_session, f"admin_{uuid.uuid4().hex[:6]}@inst.edu", role="admin", institution_id=inst.id)
        student = await create_user(db_session, f"st_{uuid.uuid4().hex[:6]}@inst.edu", role="student", institution_id=inst.id, cohort_id=cohort.id)

        admin_token = await login(client, admin.email)
        student_token = await login(client, student.email)

        # Create draft exam
        e_resp = await client.post(
            "/api/v1/admin/marks/exams",
            json={
                "cohort_id": str(cohort.id),
                "exam_type": "quiz",
                "title": "Surprise Quiz",
                "subject": "Geography",
                "max_marks": 10.0,
            },
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        exam_id = e_resp.json()["id"]

        # Student sees 0 exams while draft
        s_res1 = await client.get("/api/v1/student/marks", headers={"Authorization": f"Bearer {student_token}"})
        assert len(s_res1.json()["exams"]) == 0

        # Publish exam
        await client.post(f"/api/v1/admin/marks/exams/{exam_id}/publish", headers={"Authorization": f"Bearer {admin_token}"})

        # Student now sees 1 exam
        s_res2 = await client.get("/api/v1/student/marks", headers={"Authorization": f"Bearer {student_token}"})
        assert len(s_res2.json()["exams"]) == 1

        # Unpublish exam
        await client.post(f"/api/v1/admin/marks/exams/{exam_id}/unpublish", headers={"Authorization": f"Bearer {admin_token}"})

        # Student sees 0 exams again
        s_res3 = await client.get("/api/v1/student/marks", headers={"Authorization": f"Bearer {student_token}"})
        assert len(s_res3.json()["exams"]) == 0

    async def test_published_exam_cannot_be_hard_deleted(self, client: AsyncClient, db_session: AsyncSession):
        inst = await create_institution(db_session, "Delete Inst")
        cohort = await create_cohort(db_session, inst.id, "Grade 7")
        admin = await create_user(db_session, f"admin_{uuid.uuid4().hex[:6]}@inst.edu", role="admin", institution_id=inst.id)
        token = await login(client, admin.email)
        headers = {"Authorization": f"Bearer {token}"}

        # Create & publish
        e_resp = await client.post(
            "/api/v1/admin/marks/exams",
            json={
                "cohort_id": str(cohort.id),
                "exam_type": "internals",
                "title": "Term 1",
                "subject": "Science",
                "max_marks": 25.0,
            },
            headers=headers,
        )
        exam_id = e_resp.json()["id"]

        await client.post(f"/api/v1/admin/marks/exams/{exam_id}/publish", headers=headers)

        # Attempt hard delete on published exam -> 400
        del_resp = await client.delete(f"/api/v1/admin/marks/exams/{exam_id}", headers=headers)
        assert del_resp.status_code == 400
        assert "cannot be deleted" in del_resp.json()["detail"].lower()

    async def test_edit_after_publish_requires_reason_and_makes_change_record(self, client: AsyncClient, db_session: AsyncSession):
        inst = await create_institution(db_session, "Audit Inst")
        cohort = await create_cohort(db_session, inst.id, "Grade 10")
        admin = await create_user(db_session, f"admin_{uuid.uuid4().hex[:6]}@inst.edu", role="admin", institution_id=inst.id)
        student = await create_user(db_session, f"st_{uuid.uuid4().hex[:6]}@inst.edu", role="student", institution_id=inst.id, cohort_id=cohort.id)

        token = await login(client, admin.email)
        headers = {"Authorization": f"Bearer {token}"}

        # Create, save initial mark, publish
        e_resp = await client.post(
            "/api/v1/admin/marks/exams",
            json={
                "cohort_id": str(cohort.id),
                "exam_type": "mid-term",
                "title": "Midterm",
                "subject": "Civics",
                "max_marks": 50.0,
            },
            headers=headers,
        )
        exam_id = e_resp.json()["id"]

        await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/marks",
            json={"marks": [{"student_id": str(student.id), "marks_obtained": 40.0, "is_absent": False}]},
            headers=headers,
        )
        await client.post(f"/api/v1/admin/marks/exams/{exam_id}/publish", headers=headers)

        # Edit mark WITHOUT reason -> 400
        fail_edit = await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/marks",
            json={"marks": [{"student_id": str(student.id), "marks_obtained": 45.0, "is_absent": False}], "reason": ""},
            headers=headers,
        )
        assert fail_edit.status_code == 400
        assert "requires a reason" in fail_edit.json()["detail"].lower()

        # Edit mark WITH reason -> 200 and audit record created
        ok_edit = await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/marks",
            json={
                "marks": [{"student_id": str(student.id), "marks_obtained": 45.0, "is_absent": False}],
                "reason": "Re-evaluation of question 4",
            },
            headers=headers,
        )
        assert ok_edit.status_code == 200

        # Check MarkChangeRecord in DB
        cr = await db_session.scalar(
            select(MarkChangeRecord).where(
                MarkChangeRecord.exam_id == uuid.UUID(exam_id),
                MarkChangeRecord.student_id == student.id,
            )
        )
        assert cr is not None
        assert cr.previous_marks == Decimal("40.00")
        assert cr.new_marks == Decimal("45.00")
        assert cr.reason == "Re-evaluation of question 4"

    async def test_student_sees_only_own_marks_and_independent_gets_empty(self, client: AsyncClient, db_session: AsyncSession):
        inst = await create_institution(db_session, "Isolation Inst")
        cohort = await create_cohort(db_session, inst.id, "Grade 9-B")
        admin = await create_user(db_session, f"admin_{uuid.uuid4().hex[:6]}@inst.edu", role="admin", institution_id=inst.id)
        st_a = await create_user(db_session, f"sta_{uuid.uuid4().hex[:6]}@inst.edu", role="student", institution_id=inst.id, cohort_id=cohort.id)
        st_b = await create_user(db_session, f"stb_{uuid.uuid4().hex[:6]}@inst.edu", role="student", institution_id=inst.id, cohort_id=cohort.id)
        st_indep = await create_user(db_session, f"indep_{uuid.uuid4().hex[:6]}@indep.edu", role="student", institution_id=None, cohort_id=None)

        admin_token = await login(client, admin.email)
        token_a = await login(client, st_a.email)
        token_indep = await login(client, st_indep.email)

        # Create & save marks & publish
        e_resp = await client.post(
            "/api/v1/admin/marks/exams",
            json={
                "cohort_id": str(cohort.id),
                "exam_type": "unit-test",
                "title": "Test 1",
                "subject": "Algebra",
                "max_marks": 20.0,
            },
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        exam_id = e_resp.json()["id"]

        await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/marks",
            json={
                "marks": [
                    {"student_id": str(st_a.id), "marks_obtained": 18.0, "is_absent": False},
                    {"student_id": str(st_b.id), "marks_obtained": 12.0, "is_absent": False},
                ]
            },
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        await client.post(f"/api/v1/admin/marks/exams/{exam_id}/publish", headers={"Authorization": f"Bearer {admin_token}"})

        # Student A views marks
        resp_a = await client.get("/api/v1/student/marks", headers={"Authorization": f"Bearer {token_a}"})
        assert resp_a.status_code == 200
        data_a = resp_a.json()
        assert len(data_a["exams"]) == 1
        assert float(data_a["exams"][0]["marks_obtained"]) == 18.0
        assert float(data_a["exams"][0]["percentage"]) == 90.0
        # Ensure Student B's info is NOT in response
        assert "12" not in str(data_a)
        assert st_b.email not in str(data_a)
        assert "rank" not in data_a
        assert "class_average" not in data_a

        # Independent student views marks -> clean empty result
        resp_indep = await client.get("/api/v1/student/marks", headers={"Authorization": f"Bearer {token_indep}"})
        assert resp_indep.status_code == 200
        assert resp_indep.json()["enrolled"] is False
        assert len(resp_indep.json()["exams"]) == 0

    async def test_excel_template_preview_and_confirm_workflow(self, client: AsyncClient, db_session: AsyncSession):
        inst = await create_institution(db_session, "Excel Inst")
        cohort = await create_cohort(db_session, inst.id, "Grade 6")
        admin = await create_user(db_session, f"admin_{uuid.uuid4().hex[:6]}@inst.edu", role="admin", institution_id=inst.id)
        st1 = await create_user(db_session, f"st1_{uuid.uuid4().hex[:6]}@inst.edu", role="student", institution_id=inst.id, cohort_id=cohort.id)
        st2 = await create_user(db_session, f"st2_{uuid.uuid4().hex[:6]}@inst.edu", role="student", institution_id=inst.id, cohort_id=cohort.id)

        token = await login(client, admin.email)
        headers = {"Authorization": f"Bearer {token}"}

        # Create exam
        e_resp = await client.post(
            "/api/v1/admin/marks/exams",
            json={
                "cohort_id": str(cohort.id),
                "exam_type": "unit-test",
                "title": "Excel Test",
                "subject": "Computers",
                "max_marks": 25.0,
            },
            headers=headers,
        )
        exam_id = e_resp.json()["id"]

        # 1. Download template
        tpl_resp = await client.get(f"/api/v1/admin/marks/exams/{exam_id}/template", headers=headers)
        assert tpl_resp.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(tpl_resp.content))
        ws = wb.active
        assert ws["A1"].value == "Class:"

        # 2. Upload valid Excel sheet for preview
        # Create an Excel with valid row for st1 and an invalid row (> max marks) for st2
        wb_upload = openpyxl.Workbook()
        ws_up = wb_upload.active
        ws_up.append(["Full Name", "Email", "Username", "Max Marks", "Marks Obtained", "Absent (Y/N)"])
        ws_up.append([st1.full_name, st1.email, st1.username, 25.0, 22.5, ""])
        ws_up.append([st2.full_name, st2.email, st2.username, 25.0, 30.0, ""])  # Exceeds max 25

        buf = io.BytesIO()
        wb_upload.save(buf)
        excel_bytes = buf.getvalue()

        prev_resp = await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/preview",
            files={"file": ("marks.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers,
        )
        assert prev_resp.status_code == 200
        p_data = prev_resp.json()
        assert p_data["valid_rows_count"] == 1
        assert p_data["invalid_rows_count"] == 1
        assert "exceed maximum" in p_data["invalid_rows"][0]["error_reason"].lower()

        # 3. Download error report
        err_resp = await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/error-report",
            json=p_data["invalid_rows"],
            headers=headers,
        )
        assert err_resp.status_code == 200
        assert len(err_resp.content) > 100

        # 4. Confirm saving the valid row
        confirm_resp = await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/confirm",
            json={
                "rows": [
                    {"student_id": str(st1.id), "marks_obtained": 22.5, "is_absent": False}
                ]
            },
            headers=headers,
        )
        assert confirm_resp.status_code == 200
        assert confirm_resp.json()["created"] == 1

        # Check DB
        m = await db_session.scalar(select(Mark).where(Mark.exam_id == uuid.UUID(exam_id), Mark.student_id == st1.id))
        assert m.marks_obtained == Decimal("22.50")

    async def test_publish_unpublish_edit_return_cohort_name_no_lazy_load(
        self, client: AsyncClient, db_session: AsyncSession
    ):
        """Regression test for MissingGreenlet bug.

        Every endpoint that builds an ExamResponse (publish, unpublish,
        update, get-details) must return ``cohort_name`` and must NOT
        trigger a lazy relationship load on ``exam.cohort``.

        With the old code, publish / unpublish / update_exam crashed with
        sqlalchemy.exc.MissingGreenlet because ``_format_exam_response``
        fell back to ``exam.cohort.name`` after commit.
        """
        inst = await create_institution(db_session, "Greenlet Inst")
        cohort = await create_cohort(db_session, inst.id, "GreenletClass")
        admin = await create_user(
            db_session,
            f"admin_{uuid.uuid4().hex[:6]}@gl.edu",
            role="admin",
            institution_id=inst.id,
        )
        student = await create_user(
            db_session,
            f"st_{uuid.uuid4().hex[:6]}@gl.edu",
            role="student",
            institution_id=inst.id,
            cohort_id=cohort.id,
        )
        token = await login(client, admin.email)
        headers = {"Authorization": f"Bearer {token}"}

        # Create exam
        e_resp = await client.post(
            "/api/v1/admin/marks/exams",
            json={
                "cohort_id": str(cohort.id),
                "exam_type": "internals",
                "title": "Greenlet Test",
                "subject": "CS",
                "max_marks": 50.0,
            },
            headers=headers,
        )
        assert e_resp.status_code == 200, e_resp.text
        exam_id = e_resp.json()["id"]
        assert e_resp.json()["cohort_name"] == "GreenletClass"

        # GET details
        det_resp = await client.get(
            f"/api/v1/admin/marks/exams/{exam_id}", headers=headers
        )
        assert det_resp.status_code == 200, det_resp.text
        assert det_resp.json()["cohort_name"] == "GreenletClass"

        # PATCH (update exam title)
        upd_resp = await client.patch(
            f"/api/v1/admin/marks/exams/{exam_id}",
            json={"title": "Updated Greenlet Test"},
            headers=headers,
        )
        assert upd_resp.status_code == 200, upd_resp.text
        assert upd_resp.json()["cohort_name"] == "GreenletClass"

        # PUBLISH
        pub_resp = await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/publish",
            headers=headers,
        )
        assert pub_resp.status_code == 200, pub_resp.text
        assert pub_resp.json()["exam"]["cohort_name"] == "GreenletClass"

        # UNPUBLISH
        unpub_resp = await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/unpublish",
            headers=headers,
        )
        assert unpub_resp.status_code == 200, unpub_resp.text
        assert unpub_resp.json()["cohort_name"] == "GreenletClass"

    async def test_data_integrity_publish_and_edit_save_exactly_once(
        self, client: AsyncClient, db_session: AsyncSession
    ):
        """Confirm that a publish or edit that previously returned 500
        (due to MissingGreenlet) was saved exactly once: no half-saved
        state, and a mark edit creates exactly one MarkChangeRecord.
        """
        inst = await create_institution(db_session, "Integrity Inst")
        cohort = await create_cohort(db_session, inst.id, "IntClass")
        admin = await create_user(
            db_session,
            f"admin_{uuid.uuid4().hex[:6]}@int.edu",
            role="admin",
            institution_id=inst.id,
        )
        student = await create_user(
            db_session,
            f"st_{uuid.uuid4().hex[:6]}@int.edu",
            role="student",
            institution_id=inst.id,
            cohort_id=cohort.id,
        )
        token = await login(client, admin.email)
        headers = {"Authorization": f"Bearer {token}"}

        # Create exam
        e_resp = await client.post(
            "/api/v1/admin/marks/exams",
            json={
                "cohort_id": str(cohort.id),
                "exam_type": "end-term",
                "title": "Integrity Exam",
                "subject": "Logic",
                "max_marks": 100.0,
            },
            headers=headers,
        )
        exam_id = e_resp.json()["id"]

        # Save initial mark
        save1 = await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/marks",
            json={
                "marks": [
                    {"student_id": str(student.id), "marks_obtained": 85.0, "is_absent": False}
                ]
            },
            headers=headers,
        )
        assert save1.status_code == 200

        # Publish — must succeed and save exactly once
        pub = await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/publish", headers=headers
        )
        assert pub.status_code == 200

        # Verify exam status is published exactly once
        exam_db = await db_session.get(Exam, uuid.UUID(exam_id))
        await db_session.refresh(exam_db)
        assert exam_db.status == "published"
        assert exam_db.published_at is not None

        # Only one Mark row exists for this student+exam
        mark_count_stmt = select(func.count(Mark.id)).where(
            Mark.exam_id == uuid.UUID(exam_id),
            Mark.student_id == student.id,
        )
        mark_count = await db_session.scalar(mark_count_stmt)
        assert mark_count == 1, f"Expected 1 mark row, got {mark_count}"

        # Edit mark after publish — must create exactly 1 MarkChangeRecord
        edit = await client.post(
            f"/api/v1/admin/marks/exams/{exam_id}/marks",
            json={
                "marks": [
                    {"student_id": str(student.id), "marks_obtained": 90.0, "is_absent": False}
                ],
                "reason": "Rechecking answer sheet",
            },
            headers=headers,
        )
        assert edit.status_code == 200

        # Count change records — exactly 1
        cr_count_stmt = select(func.count(MarkChangeRecord.id)).where(
            MarkChangeRecord.exam_id == uuid.UUID(exam_id),
            MarkChangeRecord.student_id == student.id,
        )
        cr_count = await db_session.scalar(cr_count_stmt)
        assert cr_count == 1, f"Expected 1 change record, got {cr_count}"

        # Verify the change record contents
        cr = await db_session.scalar(
            select(MarkChangeRecord).where(
                MarkChangeRecord.exam_id == uuid.UUID(exam_id),
                MarkChangeRecord.student_id == student.id,
            )
        )
        assert cr.previous_marks == Decimal("85.00")
        assert cr.new_marks == Decimal("90.00")
        assert cr.reason == "Rechecking answer sheet"

        # Mark value is updated to 90, not duplicated
        mark = await db_session.scalar(
            select(Mark).where(
                Mark.exam_id == uuid.UUID(exam_id),
                Mark.student_id == student.id,
            )
        )
        assert mark.marks_obtained == Decimal("90.00")

